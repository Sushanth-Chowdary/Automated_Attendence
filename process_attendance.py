import os
import glob
import re
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def process_attendance(base_dir, target_folder):
    input_dir = os.path.join(base_dir, target_folder)
    
    # Define file paths
    true_attendance_path = os.path.join(base_dir, "filtered_true_attendance.csv")
    name_mapping_path = os.path.join(base_dir, "name_mapping.csv")
    protocol_mapping_path = os.path.join(base_dir, "protocol_mapping.csv")
    
    # 1. Check if required files exist
    for path in [true_attendance_path, name_mapping_path, protocol_mapping_path]:
        if not os.path.exists(path):
            print(f"Error: Could not find required file at {path}")
            return

    # 2. Load True Attendance and Create Roll + Name Combo
    true_df = pd.read_csv(true_attendance_path)
    
    roll_col = true_df.iloc[:, 0].astype(str).str.strip()
    name_col = true_df.iloc[:, 1].astype(str).str.strip()
    
    # Create mapping dictionary to convert Roll No to "Roll No - Student Name"
    true_df['Combo_Name'] = roll_col + " - " + name_col
    roll_to_combo = dict(zip(roll_col, true_df['Combo_Name']))
    
    # Force index to be strictly stripped Roll Numbers
    true_df = true_df.set_index(roll_col)
    
    # Clean up True Attendance headers (strip spaces and remove hyphens to match video dates)
    true_df.columns = true_df.columns.astype(str).str.strip().str.replace('-', '')
    
    # Clean cells to remove hidden spaces
    for col in true_df.columns:
        if col != 'Combo_Name':
            true_df[col] = true_df[col].astype(str).str.strip()
    
    # Load mappings
    name_map_df = pd.read_csv(name_mapping_path).dropna(subset=['Short_Name', 'Roll_Name'])
    name_map_df = name_map_df[name_map_df['Short_Name'].astype(str).str.strip() != '']
    short_to_roll = dict(zip(
        name_map_df['Short_Name'].astype(str).str.strip(),
        name_map_df['Roll_Name'].astype(str).str.strip()
    ))
    
    protocol_map_df = pd.read_csv(protocol_mapping_path)
    base_to_protocol = dict(zip(
        protocol_map_df['Base_Timestamp'].astype(str).str.strip(),
        protocol_map_df['Protocol_Length'].astype(str).str.strip()
    ))
    
    # 3. Scan for video output CSVs
    file_pattern = os.path.join(input_dir, "*_output.csv")
    output_files = glob.glob(file_pattern)
    
    if not output_files:
        print(f"No output CSV files found in {input_dir}")
        return

    # 4. Group files by base timestamp (to pair cam1 and cam2)
    sessions = {}
    for file in output_files:
        filename = os.path.basename(file)
        parts = filename.split('_')
        if len(parts) >= 2:
            base_timestamp = f"{parts[0]}_{parts[1]}"
            if base_timestamp not in sessions:
                sessions[base_timestamp] = []
            sessions[base_timestamp].append(file)

    dfs = []
    timestamps_ordered = sorted(list(sessions.keys()))

    print("Processing grouped camera files and calculating union...")
    
    # Trackers for efficiency calculations
    overall_student_matches = None
    overall_student_valid = None
    protocol_student_matches = {}
    protocol_student_valid = {}
    
    protocol_stats = {}
    total_matches_overall = 0
    total_possible_overall = 0
    
    bottom_summary_row = {}

    for ts in timestamps_ordered:
        date_str = ts.split('_')[0]
        protocol = base_to_protocol.get(ts, "Unknown Length")
        col_group = f"{ts} ({protocol})"
        
        session_files = sessions[ts]
        
        cam_dfs = [pd.read_csv(f) for f in session_files]
        merged_cams = pd.concat(cam_dfs, ignore_index=True)
        
        merged_cams['Name'] = merged_cams['Name'].astype(str).str.strip()
        merged_cams['Roll_Name'] = merged_cams['Name'].map(short_to_roll).fillna(merged_cams['Name'])
        
        merged_cams['Is_Present'] = merged_cams['Status'].astype(str).str.strip().str.lower() == 'present'
        
        grouped = merged_cams.groupby('Roll_Name').agg(
            Is_Present=('Is_Present', 'any'),
            Detection_Count=('Detection Count', 'sum') 
        ).reset_index()
        
        grouped['Estimated'] = grouped['Is_Present'].apply(lambda x: 'Present' if x else 'Absent')
        grouped = grouped.rename(columns={'Detection_Count': 'Count'})
        grouped = grouped.set_index('Roll_Name')
        
        # Fetch True Attendance for the specific date
        if date_str in true_df.columns:
            grouped['True'] = true_df[date_str]
        else:
            grouped['True'] = ""  # Leave blank if no true data exists for this date
            
        grouped = grouped[['Estimated', 'Count', 'True']]
        
        header_level_2 = col_group
        grouped.columns = pd.MultiIndex.from_product([[date_str], [header_level_2], grouped.columns])
        dfs.append(grouped)

    # Combine session data
    combined_df = pd.concat(dfs, axis=1)
    
    # Initialize trackers using the combined index
    idx = combined_df.index
    overall_student_matches = pd.Series(0.0, index=idx)
    overall_student_valid = pd.Series(0.0, index=idx)

    # 5. Calculate Efficiencies
    for ts in timestamps_ordered:
        date_str = ts.split('_')[0]
        protocol = base_to_protocol.get(ts, "Unknown Length")
        col_group = f"{ts} ({protocol})"
        
        if protocol not in protocol_stats:
            protocol_stats[protocol] = {'matches': 0, 'total': 0}
            protocol_student_matches[protocol] = pd.Series(0.0, index=idx)
            protocol_student_valid[protocol] = pd.Series(0.0, index=idx)
            
        est_col = combined_df[(date_str, col_group, 'Estimated')].astype(str).str.strip().str.lower()
        true_col = combined_df[(date_str, col_group, 'True')].astype(str).str.strip().str.lower()
        
        # Determine valid comparisons
        valid_mask = (~true_col.isin(['unknown', 'nan', 'none', ''])) & (true_col != '')
        matches = (est_col == true_col) & valid_mask
        
        # Track student-level data
        overall_student_matches += matches.astype(int)
        overall_student_valid += valid_mask.astype(int)
        protocol_student_matches[protocol] += matches.astype(int)
        protocol_student_valid[protocol] += valid_mask.astype(int)
        
        # Session-level stats for the bottom row
        current_matches = matches.sum()
        current_total = valid_mask.sum()
        
        if current_total > 0:
            date_accuracy = (current_matches / current_total) * 100
            bottom_summary_row[(date_str, col_group, 'Estimated')] = f"{date_accuracy:.2f}%"
            protocol_stats[protocol]['matches'] += current_matches
            protocol_stats[protocol]['total'] += current_total
            total_matches_overall += current_matches
            total_possible_overall += current_total
        else:
            bottom_summary_row[(date_str, col_group, 'Estimated')] = "N/A"
            
        bottom_summary_row[(date_str, col_group, 'Count')] = ""
        bottom_summary_row[(date_str, col_group, 'True')] = ""

    # Sort protocols logically (e.g., "5 min", "10 min", "15 min")
    def sort_key(p):
        m = re.search(r'\d+', p)
        return int(m.group()) if m else 999

    sorted_protocols = sorted(protocol_stats.keys(), key=sort_key)

    # Add Protocol Efficiency Columns
    for prot in sorted_protocols:
        col_name = (f'{prot} Protocol', 'Summary', 'Efficiency')
        mask = protocol_student_valid[prot] > 0
        
        prot_series = pd.Series(pd.NA, index=idx)
        prot_series[mask] = (protocol_student_matches[prot][mask] / protocol_student_valid[prot][mask]) * 100
        
        combined_df[col_name] = prot_series.apply(
            lambda x: f"{x:.2f}%" if pd.notna(x) else "N/A"
        )
        
        # Bottom row for protocol column
        p_matches = protocol_stats[prot]['matches']
        p_total = protocol_stats[prot]['total']
        bottom_summary_row[col_name] = f"{(p_matches / p_total) * 100:.2f}%" if p_total > 0 else "N/A"

    # Add Overall Efficiency Column
    overall_col = ('Overall', 'Summary', 'Efficiency')
    mask_overall = overall_student_valid > 0
    overall_series = pd.Series(pd.NA, index=idx)
    overall_series[mask_overall] = (overall_student_matches[mask_overall] / overall_student_valid[mask_overall]) * 100
    
    combined_df[overall_col] = overall_series.apply(
        lambda x: f"{x:.2f}%" if pd.notna(x) else "N/A"
    )
    
    bottom_summary_row[overall_col] = f"{(total_matches_overall / total_possible_overall) * 100:.2f}%" if total_possible_overall > 0 else "N/A"

    # Convert Index to Full Combo Name and sort
    combined_df.index = combined_df.index.map(lambda x: roll_to_combo.get(str(x), str(x)))
    combined_df.index.name = 'Student Name & Roll No'
    combined_df = combined_df.sort_index()

    # Append Summary Row
    summary_df = pd.DataFrame([bottom_summary_row], index=['AVERAGE EFFICIENCY'])
    combined_df = pd.concat([combined_df, summary_df])

    # 8. Save Files
    csv_output = os.path.join(input_dir, "Final_Report_With_Efficiency.csv")
    xlsx_output = os.path.join(input_dir, "Final_Report_With_Efficiency.xlsx")

    print("Saving merged CSV and formatting Excel report...")
    combined_df.to_csv(csv_output)
    combined_df.to_excel(xlsx_output)

    # 9. Excel Formatting
    wb = openpyxl.load_workbook(xlsx_output)
    ws = wb.active
    ws.title = "Attendance Efficiency"

    ws['A1'] = 'Student Name & Roll No'
    ws.merge_cells('A1:A3')
    ws.delete_rows(4)  
    
    max_row = ws.max_row
    max_col = ws.max_column

    header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    green_font = Font(color="006100")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    red_font = Font(color="9C0006")
    summary_fill = PatternFill(start_color="D5F5E3", end_color="D5F5E3", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(left=Side(style='thin', color='BDC3C7'), right=Side(style='thin', color='BDC3C7'), 
                         top=Side(style='thin', color='BDC3C7'), bottom=Side(style='thin', color='BDC3C7'))

    # Format Headers (Rows 1, 2, and 3)
    for row in range(1, 4):
        for cell in ws[row]:
            try:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align
                cell.border = thin_border
            except AttributeError:
                pass

    # Format Data Rows (up to the last row, which is the summary)
    for row_idx in range(4, max_row):
        for col_idx, cell in enumerate(ws[row_idx], start=1):
            try:
                cell.alignment = left_align if col_idx == 1 else center_align
                cell.border = thin_border
                
                if row_idx % 2 == 0:
                    cell.fill = PatternFill(start_color="F8F9F9", end_color="F8F9F9", fill_type="solid")
                    
                val = str(cell.value).strip().lower()
                if val == 'present':
                    cell.fill = green_fill
                    cell.font = green_font
                elif val == 'absent':
                    cell.fill = red_fill
                    cell.font = red_font
                elif val.endswith('%'):
                    if float(val.strip('%')) < 80.0:
                        cell.fill = red_fill
                        cell.font = red_font
            except (AttributeError, ValueError):
                pass

    # Format the single Summary Row at the bottom
    for col_idx, cell in enumerate(ws[max_row], start=1):
        try:
            cell.alignment = center_align if col_idx > 1 else left_align
            cell.border = thin_border
            cell.fill = summary_fill
            cell.font = Font(bold=True)
            
            val = str(cell.value).strip().lower()
            if val.endswith('%') and val != 'n/a':
                if float(val.strip('%')) < 80.0:
                    cell.fill = red_fill
                    cell.font = Font(color="9C0006", bold=True)
        except (AttributeError, ValueError):
            pass

    # Merge Estimated, Count, True columns for ONLY the session columns in the summary row
    for i in range(len(timestamps_ordered)):
        start_col = 2 + (i * 3)
        end_col = start_col + 2
        ws.merge_cells(start_row=max_row, start_column=start_col, end_row=max_row, end_column=end_col)

    # Adjust column widths
    for col_idx in range(1, max_col + 1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = 15
    ws.column_dimensions['A'].width = 38

    ws.freeze_panes = "B4"
    wb.save(xlsx_output)
    
    print(f"\nSuccess! Reports generated:")
    print(f" -> {csv_output}")
    print(f" -> {xlsx_output}")

if __name__ == "__main__":
    BASE_DIRECTORY = "ATTENDENCE RESULTS/Results"
    TARGET_FOLDER = "test 1"  
    
    process_attendance(BASE_DIRECTORY, TARGET_FOLDER)